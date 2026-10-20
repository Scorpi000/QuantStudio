"""生成 QuantStudio 系统介绍演示文稿及预览。

通过同一套场景描述生成原生可编辑 PowerPoint、矢量 PDF 和 PNG 预览。
核心介绍依据三个系统项目的本地源码与文档，Barra 因子示例依据 QSResearch
中的定义脚本；每页备注保存讲稿与来源。
脚本不连接业务数据库、不执行研究策略，也不修改现有业务代码。

使用方法：在项目根目录使用 QS 虚拟环境运行本脚本。可选先运行
``node scripts/capture_quantstudio_presentation.mjs`` 生成真实工作台演示截图。

Args:
    --output: 输出目录，默认 docs/presentation。

输出：中文 PPTX、PDF、总览图、逐页预览、来源清单与版式校验结果。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import subprocess
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
W, H = 1600, 900
BG = '#0A1324'
CARD = '#111F34'
CARD2 = '#162941'
WHITE = '#F1F5FA'
MUTED = '#9CAEC5'
DIM = '#59718E'
LINE = '#29405B'
TEAL = '#65E0C4'
BLUE = '#81B5FF'
GOLD = '#DCC49A'
FONT_FILE = 'C:/Windows/Fonts/msyh.ttc'
BOLD_FILE = 'C:/Windows/Fonts/msyhbd.ttc'
LATIN_FILE = 'C:/Windows/Fonts/segoeui.ttf'
LATIN_BOLD = 'C:/Windows/Fonts/seguisb.ttf'

SOURCES = {
    'Q01': ('QuantStudio', 'docs/通则和约定.ipynb'),
    'Q02': ('QuantStudio', 'docs/Core/计算图框架.ipynb'),
    'Q03': ('QuantStudio', 'QuantStudio/Core/Node.py'),
    'Q04': ('QuantStudio', 'QuantStudio/Core/CalcEngine.py'),
    'Q05': ('QuantStudio', 'QuantStudio/Core/ParallelEngine.py'),
    'Q06': ('QuantStudio', 'docs/因子框架/基本框架.ipynb'),
    'Q07': ('QuantStudio', 'QuantStudio/Factor/api.py'),
    'Q08': ('QuantStudio', 'QuantStudio/BackTest/SectionFactor/api.py'),
    'Q09': ('QuantStudio', 'docs/回测框架/策略回测.ipynb'),
    'Q10': ('QuantStudio', 'QuantStudio/BackTest/Strategy/Strategy.py'),
    'Q11': ('QuantStudio', 'docs/风险模型/基本框架.ipynb'),
    'Q12': ('QuantStudio', 'QuantStudio/Risk/RiskModel/BarraModel.py'),
    'Q13': ('QuantStudio', 'docs/组合优化/基本框架.ipynb'),
    'Q14': ('QuantStudio', 'QuantStudio/PortfolioConstructor/api.py'),
    'Q15': ('QuantStudio', 'QuantStudio/BackTest/PerformanceAnalysis/BrinsonModel.py'),
    'Q16': ('QuantStudio', 'setup.py'),
    'Q17': ('QuantStudio', 'QuantStudio/Factor/FactorOperator.py'),
    'E01': ('QSExt', 'docs/Graph/README.md'),
    'E02': ('QSExt', 'QSExt/Graph/Compiler.py'),
    'E03': ('QSExt', 'QSExt/Graph/api.py'),
    'E04': ('QSExt', 'docs/DefModule/DefModule.md'),
    'E05': ('QSExt', 'QSExt/DefModule/DefContent.py'),
    'E06': ('QSExt', 'QSExt/RuntimeConfig/config.py'),
    'E07': ('QSExt', 'QSExt/QSRegistry/QSGraphDB.py'),
    'E08': ('QSExt', 'QSExt/ReportGenerator/__init__.py'),
    'E09': ('QSExt', 'QSExt/ReportGenerator/scenarios/__init__.py'),
    'E10': ('QSExt', 'QSExt/GPFactor/GPLearn.py'),
    'E11': ('QSExt', 'QSExt/GPFactor/fitness.py'),
    'E12': ('QSExt', 'QSExt/LLMFactor/loop/round_runner.py'),
    'E13': ('QSExt', 'mcp/qs_registry.py'),
    'E14': ('QSExt', 'QSExt/Factor/AKShareDB.py'),
    'E15': ('QSExt', 'QSExt/Factor/ClickHouseDB.py'),
    'E16': ('QSExt', 'QSExt/Factor/DuckDB.py'),
    'E17': ('QSExt', 'QSExt/Factor/MongoDB.py'),
    'E18': ('QSExt', 'QSExt/Factor/Neo4jDB.py'),
    'E19': ('QSExt', 'QSExt/Factor/ZarrDB.py'),
    'E20': ('QSExt', 'QSExt/Graph/Subgraphs.py'),
    'W01': ('QSWeb', 'docs/13-统一计算图工作台.md'),
    'W02': ('QSWeb', 'frontend/src/App.tsx'),
    'W03': ('QSWeb', 'backend/app/services/graph_service.py'),
    'W04': ('QSWeb', 'backend/app/services/graph_runtime.py'),
    'W05': ('QSWeb', 'backend/app/services/graph_repository.py'),
    'W06': ('QSWeb', 'backend/app/tasks/manager.py'),
    'W07': ('QSWeb', 'backend/app/services/risk_service.py'),
    'W08': ('QSWeb', 'backend/app/services/portfolio_service.py'),
    'W09': ('QSWeb', 'backend/app/services/report_service.py'),
    'W10': ('QSWeb', 'backend/app/services/mining_service.py'),
    'W11': ('QSWeb', 'backend/app/services/ai_service.py'),
    'W12': ('QSWeb', 'backend/app/core/security.py'),
    'W13': ('QSWeb', 'backend/app/core/user_workspace.py'),
    'W14': ('QSWeb', 'docker-compose.yml'),
    'W15': ('QSWeb', 'frontend/src/components/Layout/MainLayout.tsx'),
    'W16': ('QSWeb', 'frontend/src/components/GraphEditor/index.tsx'),
    'W17': ('QSWeb', 'tests/graph_editor.mount.tsx'),
    'W18': ('QSWeb', 'frontend/src/pages/MiningStudio/index.tsx'),
    'W19': ('QSWeb', 'frontend/src/pages/Notebook/index.tsx'),
    'W20': ('QSWeb', 'docs/12-服务启动目录与初始化.md'),
    'W21': ('QSWeb', 'frontend/src/pages/PortfolioOptimizer/ObjectiveConfig.tsx'),
    'R01': ('QSResearch', 'QSResearch/FactorDef/JY/stock_cn_factor_barra.py'),
    'R02': ('QSResearch', 'QSResearch/FactorDef/JY/stock_cn_factor_barra_descriptor.py'),
}


def rgb(c):
    return tuple(bytes.fromhex(c.lstrip('#')))


def font(size, bold=False, latin=False):
    path = (LATIN_BOLD if bold else LATIN_FILE) if latin else (BOLD_FILE if bold else FONT_FILE)
    return ImageFont.truetype(path, round(size))


def wrap(text, width, size, bold=False, latin=False):
    """按字体度量断行，显式换行保留。"""
    f = font(size, bold, latin)
    lines = []
    for paragraph in text.split('\n'):
        current = ''
        for char in paragraph:
            if current and f.getlength(current + char) > width:
                lines.append(current.rstrip())
                current = char.lstrip()
            else:
                current += char
        lines.append(current)
    return lines


class Slide:
    """保存一页原生图形、讲稿和引用。"""

    def __init__(self, chapter, title, subtitle='', sources=(), notes=''):
        self.chapter, self.title, self.sources, self.notes = chapter, title, list(sources), notes
        self.items = []
        self.rect(0, 0, W, H, BG)
        self.text(80, 40, 'QUANTSTUDIO', 19, TEAL, bold=True, latin=True)
        self.text(1180, 40, chapter, 18, MUTED, width=340, align='right')
        self.text(80, 103, title, 51, WHITE, bold=True, width=1450)
        if subtitle:
            self.text(82, 185, subtitle, 25, MUTED, width=1430)

    def rect(self, x, y, w, h, fill=CARD, stroke=None, radius=0):
        self.items.append(dict(kind='rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, radius=radius))

    def text(self, x, y, value, size=28, color=WHITE, width=1400, bold=False, latin=False, align='left', leading=1.45):
        lines = wrap(value, width, size, bold, latin)
        for i, line in enumerate(lines):
            actual = font(size, bold, latin).getlength(line)
            px = x + (width - actual if align == 'right' else (width - actual) / 2 if align == 'center' else 0)
            self.items.append(dict(kind='text', x=px, y=y + i * size * leading,
                                   w=max(actual + 12, 20), h=size * 1.5, text=line,
                                   size=size, color=color, bold=bold, latin=latin))
        return len(lines) * size * leading

    def line(self, x1, y1, x2, y2, color=LINE, width=2, arrow=False):
        self.items.append(dict(kind='line', x=x1, y=y1, x2=x2, y2=y2, color=color, width=width, arrow=arrow))

    def dot(self, x, y, r=5, fill=TEAL):
        self.items.append(dict(kind='circle', x=x-r, y=y-r, w=r*2, h=r*2, fill=fill))

    def image(self, x, y, w, h, path):
        self.items.append(dict(kind='image', x=x, y=y, w=w, h=h, path=str(path)))

    def card(self, x, y, w, h, number, title, body, accent=TEAL):
        self.rect(x, y, w, h, CARD)
        self.rect(x, y, 4, h, accent)
        self.text(x+27, y+23, number, 24, accent, latin=True)
        self.text(x+27, y+76, title, 33, WHITE, bold=True, width=w-54)
        self.text(x+27, y+143, body, 27, MUTED, width=w-54)

    def node(self, x, y, w, title, subtitle='', color=TEAL, h=112):
        self.rect(x, y, w, h, CARD, LINE, 8)
        self.rect(x, y, 4, h, color)
        self.text(x+22, y+19, title, 28, WHITE, width=w-40, bold=True)
        if subtitle:
            self.text(x+22, y+66, subtitle, 19, MUTED, width=w-40, latin=subtitle.isascii())

    def footer(self, i, total):
        self.line(80, 827, 1520, 827, LINE, 1)
        self.text(80, 847, '系统介绍  /  2026.10', 15, DIM)
        self.text(400, 847, '依据：' + ' · '.join(self.sources), 14, DIM, width=1020, align='right')
        self.text(1440, 840, f'{i:02d} / {total:02d}', 20, MUTED, width=80, align='right', latin=True)


def deck(out):
    """构建二十一页系统介绍。"""
    pages = []
    def add(*args, **kwargs):
        s = Slide(*args, **kwargs)
        pages.append(s)
        return s

    s = add('系统总览', '', sources=['Q02','E01','W01'], notes='开场建议：用一句话定义系统。QuantStudio 系统由计算框架、扩展能力与 Web 工作台共同组成，目标是把数据、模型、研究验证和成果沉淀连接起来。介绍依据为 2026 年 10 月 20 日的本地工作区，包含未提交改动；不是发布版本或生产验收证明。')
    s.text(80, 190, 'QuantStudio', 100, WHITE, bold=True, latin=True, width=900)
    s.text(86, 340, '让量化研究成为\n可积累的系统能力', 58, WHITE, width=880, bold=True, leading=1.35)
    s.text(86, 535, '量化计算框架 × 研究扩展生态 × 统一 Web 工作台', 27, MUTED, width=1050)
    s.text(86, 742, 'QUANTSTUDIO  /  QSEXT  /  QSWEB', 20, MUTED, latin=True)
    # 封面图形表达计算依赖网络，全部为可编辑矢量。
    pts = [(1110,225),(1310,185),(1460,320),(1180,410),(1390,515),(1110,625),(1470,700)]
    for a,b in [(0,1),(1,2),(0,3),(2,3),(2,4),(3,4),(3,5),(4,6),(5,6)]:
        s.line(*pts[a], *pts[b], LINE, 3)
    for i,(x,y) in enumerate(pts):
        s.dot(x,y,12,TEAL if i in (0,3,5) else BLUE)
        s.dot(x,y,4,BG)
    s.text(1140, 750, '从计算到知识的连接', 20, MUTED, width=360, align='right')

    s = add('01  /  系统定位', '一个系统，贯通量化研究的三个关键环节',
            '从模型表达、研究验证，到可检索与可复用的研究资产。',
            ['Q02','Q06','E07','W01'], '本页是管理层摘要。第一，统一计算表达：因子、回测等计算对象建立在 Node 和引擎之上。第二，研究链路贯通：数据、信号、账户、风险和报告可以在同一体系组织。第三，研究资产沉淀：QSRegistry 管理计算对象元信息和依赖关系，QSWeb 保存图文档修订与运行记录。这里的价值是架构带来的组织能力，未声明量化的提效或收益。')
    s.card(80,280,460,390,'01','统一计算表达','以计算图组织因子、策略、\n回测与报告，减少重复编排。')
    s.card(570,280,460,390,'02','贯通研究验证','把数据加工、因子检验、\n账户模拟与组合约束串起来。', BLUE)
    s.card(1060,280,460,390,'03','沉淀研究资产','保存定义、依赖、修订与结果，\n支持检索、追溯及后续复用。', GOLD)
    s.text(80,727,'管理视角：把个人研究成果转化为团队可以持续使用的基础能力。',30,WHITE,width=1440)

    s = add('01  /  系统定位', '把分散的研究动作，组织成连续的研究链路',
            '系统设计回应的是数据、模型与结果之间的衔接问题。',
            ['Q06','E04','E07','W01'], '以下左列是量化研究常见的组织问题，用于说明系统的设计动机，不是针对当前团队的诊断结论。右列对应可在代码中确认的机制。收益方向是减少重复开发、提高可追溯性；实际效果应通过试点测量。')
    s.text(108,270,'常见研究断点',23,MUTED)
    s.text(795,270,'QuantStudio 的对应机制',23,TEAL)
    rows=[('数据接口各异','统一因子库 / 因子表 / 因子接口'),('脚本之间靠人工衔接','依赖声明与计算图编排'),('计算结果与报告分离','计算节点与报告节点连接'),('成果难以检索和复用','注册中心、血缘查询与图修订')]
    for i,(left,right) in enumerate(rows):
        y=326+i*107
        s.rect(80,y,1440,89,CARD)
        s.text(108,y+24,left,29,WHITE,width=560)
        s.line(656,y+44,731,y+44,DIM,2,True)
        s.text(795,y+24,right,29,WHITE,width=690)

    s = add('02  /  总体架构', '三层分工，让计算内核与应用界面各司其职',
            'QuantStudio 管计算语义，QSExt 管扩展与图协议，QSWeb 管工作台与运行交互。',
            ['Q02','E01','W01','W03'], '从上向下讲解。QSWeb 负责页面、用户资源、图文档修订、任务与结果。QSExt Graph 接受 GraphDocument，通过定义目录、静态校验与编译构造 QuantStudio 对象，模块自身不依赖 FastAPI 或任务管理器。QSExt 的 QSRegistry、DefModule、ReportGenerator 是并列扩展能力。QuantStudio 引擎执行节点，Factor、BackTest、Risk 与 PortfolioConstructor 提供领域能力。下方两类存储区分业务数据与元信息，不把 Neo4j 描述为行情或所有结果的统一存储。')
    layers=[(270,'QSWeb','统一计算图工作台 · 数据 / 风险 / 组合 / 报告 · AI / Notebook',BLUE),
            (425,'QSExt','Graph 图协议与编译 · DefModule · QSRegistry · 报告与挖掘',TEAL),
            (580,'QuantStudio','Node / Engine · Factor · BackTest · Risk · PortfolioConstructor',GOLD)]
    for y,name,body,c in layers:
        s.rect(80,y,1440,126,CARD)
        s.rect(80,y,7,126,c)
        s.text(112,y+27,name,35,c,bold=True,latin=True,width=330)
        s.text(460,y+39,body,27,WHITE,width=1010)
    s.line(814,397,814,422,DIM,2,True)
    s.line(814,552,814,577,DIM,2,True)
    s.text(93,748,'业务数据：因子库 / 风险库 / 结果文件',23,MUTED,width=700)
    s.text(840,748,'元信息与修订：Neo4j / 应用数据库',23,MUTED,width=680)

    s = add('02  /  研究链路', '从数据到决策，六个环节共享一套研究语言',
            '研究链路可围绕相同的因子、计算节点和结果组织。',
            ['Q06','Q08','Q09','Q13','E08','W01'], '此处为逻辑能力链路，并不意味着所有项目自动执行这六步。单因子研究可以从数据到检验直接生成报告；策略研究可额外使用账户回测与组合优化。底部强调沉淀与复用，由不同持久化机制共同支持，而非每次运行都自动登记到 QSRegistry。')
    names=['数据接入','因子构建','研究检验','策略回测','风险与组合','报告沉淀']
    desc=['统一读取口径','组合算子逻辑','评价预测能力','模拟账户行为','约束持仓风险','保留研究证据']
    for i,(title,body) in enumerate(zip(names,desc)):
        x=80+i*244
        s.text(x,293,f'{i+1:02d}',49,TEAL if i<3 else BLUE,latin=True)
        s.rect(x,374,220,160,CARD)
        s.text(x+16,402,title,29,WHITE,width=195,bold=True)
        s.text(x+16,464,body,21,MUTED,width=195)
        if i<5:s.line(x+220,451,x+240,451,DIM,2,True)
    s.line(190,614,1410,614,TEAL,2)
    for i in range(6):s.dot(190+i*244,614,5,TEAL)
    s.text(240,667,'定义 · 参数 · 依赖 · 图修订 · 运行记录 · 研究报告',31,WHITE,width=1120,align='center')
    s.text(240,724,'形成下一轮研究可以复用的基础',24,MUTED,width=1120,align='center')

    s = add('03  /  数据与因子', '异构数据源，进入统一的因子数据模型',
            '数据接入的差异封装在适配器内，研究逻辑围绕统一接口展开。',
            ['Q06','Q07','E14','E15','E16','E17','E18','E19'], '核心框架包含聚源、SQL、HDF5 等数据源；QSExt 扩展 AKShare、ClickHouse、DuckDB、MongoDB、Neo4j、Zarr 等适配器。适配器存在不等于所有接口都在同一部署中可用，实际需安装依赖和配置连接。右侧说明逻辑数据组织：因子表为因子 × 时点 × 证券，单因子为时点 × 证券的 DataFrame。数据读写权限由具体适配器决定。')
    groups=[('金融数据','聚源 JYDB · AKShare'),('关系与分析存储','SQL · ClickHouse · DuckDB'),('文件与专用后端','HDF5 · Zarr · MongoDB · Neo4j')]
    for i,(t,b) in enumerate(groups):
        s.rect(80,278+i*145,665,121,CARD)
        s.text(106,297+i*145,t,25,TEAL,bold=True,width=590)
        s.text(106,344+i*145,b,26,WHITE,width=590,latin=False)
    s.line(770,489,850,489,DIM,3,True)
    s.text(906,285,'FactorDB → FactorTable → Factor',26,TEAL,latin=True,width=610)
    for i in range(3):
        s.rect(978+i*25,411-i*20,390,184,CARD2,LINE)
    for r in range(1,5):s.line(1028,371+r*37,1418,371+r*37,LINE,1)
    for c in range(1,6):s.line(1028+c*65,371,1028+c*65,555,LINE,1)
    s.text(930,635,'因子 × 时点 × 证券',36,WHITE,width=540,bold=True,align='center')
    s.text(870,709,'单因子：时间索引 × 证券列的 DataFrame',23,MUTED,width=640,align='center')
    s.text(82,759,'适配器按需配置；数据授权与实际可用性由具体数据源决定。',21,DIM,width=1360)

    s = add('03  /  计算内核', '计算图，让研究逻辑具备可组合的结构',
            '节点表达计算单元，连接表达依赖，引擎负责验证与执行。',
            ['Q02','Q03','Q04','Q05'], '图中箭头表示结果流向，是对真实模型的简化示意。数据读取和衍生因子共同形成计算图，因子检验与账户回测可共享因子结果，报告依赖上游计算输出。引擎包括顺序、显式栈、多进程与树形等变体，本页仅概括单机顺序和并行能力，不宣称分布式集群或未经测量的性能。数据版本与执行键决定缓存隔离，不能仅凭同一 QSID 承诺跨运行复用。')
    s.node(80,356,245,'行情因子','读取外部数据')
    s.node(401,356,255,'衍生因子','时序 / 截面加工')
    s.node(755,278,270,'因子检验','IC / 分组分析',BLUE)
    s.node(755,484,270,'策略账户','交易 / 持仓模拟',BLUE)
    s.node(1180,356,340,'研究报告','汇聚分析结果',GOLD)
    for a in [(325,412,401,412),(656,412,755,334),(656,412,755,540),(1025,334,1180,412),(1025,540,1180,412)]:s.line(*a,TEAL,2,True)
    for i,(t,b) in enumerate([('统一生命周期','初始化 → 数据准备 → 计算'),('显式依赖与校验','命名输入 · 类型校验 · 循环检查'),('多种执行方式','顺序计算 · I/O 并发 · 多进程')]):
        x=80+i*490
        s.text(x,677,t,27,WHITE,width=460,bold=True)
        s.text(x,730,b,22,MUTED,width=460)

    s = add('03  /  模块化构建', '真实 Barra 示例：描述子复用，组合形成风格因子',
            '以 Liquidity 与 Size 为例：参数化描述子、统一预处理，以及跨因子的正交化依赖。',
            ['R01','R02'], '本页从用户指定的两份 QSResearch 脚本中节选 Liquidity 与 Size 的真实构建链路，不执行脚本或连接数据库。描述子脚本 stock_cn_factor_barra_descriptor.py 第 436–445 行：STOM 由日换手率和上市、交易标志计算，默认使用 21 日窗口并检查有效交易比例；STOQ 调用 calcSTOQ(STOM)，默认 T=3、T1=21、非空数=2；STOA 使用 calcSTOQ.new 修改 T=12、T1=21、非空数=4 及 LookBack，仍以 STOM 为输入。该算子对抽样 STOM 先取指数、平均，再取对数，不是直接算术平均对数值。第 402 行：LNCAP 为总市值的自然对数。风格因子脚本 stock_cn_factor_barra.py 第 157–160 行：每项描述子都经过 standardize → winsorize → standardize，图中星号表示处理后的描述子。standardize 使用 ESTU 中有效样本计算市值加权均值与非加权标准差，winsorize 使用该脚本定义的极值压缩方法。第 164、170、176–183 行：Size 来自处理后的 LNCAP，并在风格层标准化；Liquidity 以 STOM*、STOQ*、STOA* 按 0.35、0.35、0.30 组合，Mean 使用 ignore_nan_weight=True，然后先标准化、对 Size 回归取残差，再标准化。正交化使用的 Size 是当时的中间状态，尚未执行最终填缺处理。第 197–199 行：各风格因子再按行业和对数市值等信息填充缺失值，并最终标准化。为突出主依赖，图中折叠了市值、ESTU、行业及上市状态通向各算子的辅助连线，也未展开 Size 最后的填缺和标准化输出分支。图中模块对应实际计算步骤，不宣称原脚本已经封装为 QSExt Graph 子图组件。')

    def factor_block(x, y, w, h, title, subtitle='', accent=TEAL):
        """绘制真实因子链路中的可编辑节点。"""
        s.rect(x,y,w,h,CARD2,LINE,6)
        s.rect(x,y,3,h,accent)
        s.text(x+12,y+10,title,24,WHITE,width=w-24,bold=True)
        if subtitle:
            s.text(x+12,y+47,subtitle,19,MUTED,width=w-24)

    s.text(80,255,'规模链路：描述子 → 风格因子',22,BLUE,width=650)
    s.text(80,432,'流动性链路：共享描述子 → 参数复用 → 统一加工 → 因子组合',22,TEAL,width=1100)

    # 所有连接先于节点绘制，保留实际依赖与处理顺序。
    for x1,x2 in [(290,345),(565,640),(1000,1050)]:
        s.line(x1,342,x2,342,BLUE,2,True)
    s.line(1260,342,1425,342,BLUE,2)
    s.line(1425,342,1425,520,BLUE,2,True)
    s.text(1440,419,'解释因子',18,BLUE,width=90)

    s.line(290,530,315,530,TEAL,2)
    s.line(315,530,315,574,TEAL,2)
    s.line(315,574,340,574,TEAL,2,True)
    s.line(290,660,315,660,TEAL,2)
    s.line(315,660,315,606,TEAL,2)
    s.line(315,606,340,606,TEAL,2,True)
    s.line(540,587,570,587,TEAL,2)
    s.line(570,475,570,690,TEAL,2)
    s.dot(570,587,4,TEAL)
    s.line(570,475,850,475,TEAL,2,True)
    s.line(570,570,610,570,TEAL,2,True)
    s.line(570,690,610,690,GOLD,2,True)
    s.line(790,570,820,570,TEAL,2)
    s.line(820,570,820,585,TEAL,2)
    s.line(820,585,850,585,TEAL,2,True)
    s.line(790,690,820,690,GOLD,2)
    s.line(820,690,820,695,GOLD,2)
    s.line(820,695,850,695,GOLD,2,True)

    for y,lane,end_y,accent in [(475,1085,550,TEAL),(585,1085,585,TEAL),(695,1095,620,GOLD)]:
        s.line(1070,y,lane,y,accent,2)
        if y!=end_y:
            s.line(lane,y,lane,end_y,accent,2)
        s.line(lane,end_y,1110,end_y,accent,2,True)
    s.line(1290,589,1330,589,TEAL,2,True)
    s.line(1425,612,1425,650,TEAL,2,True)

    factor_block(80,300,210,85,'总市值','total_cap',BLUE)
    factor_block(345,300,220,85,'LNCAP','对数总市值',BLUE)
    s.rect(640,295,360,95,CARD,BLUE,6)
    s.text(654,306,'统一描述子预处理',24,WHITE,width=330,bold=True)
    s.text(654,350,'标准化 → 去极值 → 标准化',21,BLUE,width=330)
    factor_block(1050,300,210,85,'Size','风格层标准化后',BLUE)

    factor_block(80,490,210,80,'日换手率','turnover',TEAL)
    factor_block(80,620,210,80,'有效交易标志','上市 × 交易状态',TEAL)
    factor_block(340,540,200,95,'STOM','21 日换手率对数',TEAL)
    factor_block(610,530,180,80,'STOQ','3 个月 · T=3',TEAL)
    factor_block(610,650,180,80,'STOA','12 个月 · T=12',GOLD)
    s.text(610,483,'直接复用 STOM',20,TEAL,width=215)
    s.text(610,620,'同一算子，不同参数',19,GOLD,width=220)
    s.text(850,401,'同一处理链 × 三个描述子',20,TEAL,width=280)
    for y,name,accent in [(435,'STOM*',TEAL),(545,'STOQ*',TEAL),(655,'STOA*',GOLD)]:
        s.rect(850,y,220,80,CARD,accent,6)
        s.text(863,y+10,'标准化 → 去极值',22,WHITE,width=195,bold=True)
        s.text(863,y+46,f'→ 标准化 · {name}',20,accent,width=195)

    s.rect(1110,515,180,148,CARD2,LINE,6)
    s.text(1122,528,'加权合成',24,WHITE,width=156,bold=True)
    for i,label in enumerate(['STOM* × 0.35','STOQ* × 0.35','STOA* × 0.30']):
        s.text(1122,565+i*25,label,18,MUTED,width=156)
    s.text(1122,638,'→ 标准化',18,TEAL,width=156)
    factor_block(1330,520,190,92,'正交化','剔除 Size 暴露',BLUE)
    s.rect(1330,650,190,112,CARD2,TEAL,6)
    s.text(1342,659,'Liquidity',25,WHITE,width=166,bold=True,latin=True)
    s.text(1342,699,'再标准化 → 补缺\n→ 最终标准化',19,TEAL,width=166)
    s.text(82,752,'辅助口径：市值与 ESTU 用于标准化；行业、市值与上市状态用于填补缺失。',22,MUTED,width=1215)
    s.text(82,791,'真实代码节选；* 表示预处理后的描述子。辅助输入连线与 Size 最终输出分支已折叠。',18,DIM,width=1440)

    s = add('03  /  因子构建', '四类运算，覆盖点、时间、截面与面板',
            '标准算子可以嵌套组合，自定义逻辑通过明确的输入与参数声明接入。',
            ['Q06','Q07','E04'], '四类算子按依赖数据的范围区分：单点运算处理同一时间同一证券，时序运算使用历史时间序列，截面运算使用同一时点的证券集合，面板运算同时跨时间与证券。右侧例子仅用于解释范围，不意味着系统内置所有对应业务因子。自定义算子及参数声明使其可参与计算图，也可成为 Web 组件的元信息来源。')
    cases=[('单点运算','同一时点 · 同一证券','价格比值 / 条件表达式',TEAL),('时序运算','历史区间 · 同一证券','滚动均值 / 历史波动',BLUE),('截面运算','同一时点 · 全体证券','排序 / 标准化',GOLD),('面板运算','历史区间 · 全体证券','时序与截面联合计算',TEAL)]
    for i,(t,scope,ex,c) in enumerate(cases):
        x=80+(i%2)*735;y=270+(i//2)*248
        s.rect(x,y,705,220,CARD)
        s.text(x+26,y+22,t,32,c,bold=True,width=630)
        s.text(x+26,y+84,scope,27,WHITE,width=630)
        s.text(x+26,y+147,ex,25,MUTED,width=630)

    s = add('04  /  研究验证', '因子好不好，需要一组互补的检验',
            '从预测能力、组合表现、稳定性与归因四个角度形成研究证据。',
            ['Q08','Q09','Q15','W01'], '截面因子模块对外提供 IC、分位数组合、相关性、因子换手和 Fama–MacBeth 回归等分析。绩效归因用于解释收益来源，现有业绩归因模块包含配置与选择的收益分解。六个卡片对应不同的研究问题，不能由一个指标代替全部验证。PPT 不展示虚构的收益、IC 或夏普值。研究结论还需独立的样本外检验和适当的数据时间口径，这是研究流程建议。')
    tests=[('IC / 衰减','因子是否有预测信息？'),('分位数组合','排序能否转化为组合差异？'),('截面相关性','是否重复已有因子信息？'),('因子换手','信号的稳定性与实现成本？'),('Fama–MacBeth','横截面收益解释是否稳定？'),('绩效归因','收益由哪些来源贡献？')]
    for i,(t,b) in enumerate(tests):
        x=80+(i%3)*490;y=276+(i//3)*246
        s.rect(x,y,460,213,CARD)
        s.text(x+25,y+25,f'{i+1:02d}',23,TEAL,latin=True)
        s.text(x+25,y+76,t,31,WHITE,bold=True,width=410)
        s.text(x+25,y+139,b,24,MUTED,width=410)

    s = add('04  /  策略回测', '从信号到账户，再到可解释的绩效',
            '简单策略可直接组合因子，复杂策略可通过 Python 扩展。',
            ['Q09','Q10','W01'], 'QuantStudio 提供两类方式：信号因子加 MakeAccount，以及继承 MakeStrategy 并实现 genSignal。Web Python 策略遵循 build_strategy(args, factors, price) 显式入口契约。MakeAccount 接受价格、交易信号、买卖费率、买卖限制及金额限制等输入，输出现金、持仓、成交量、成交价和费用等账户数据。AccountStats 计算绩效，再由 AccountReport 输出报告。这里描述历史研究和模拟能力，不包含券商接入或实盘执行承诺。')
    s.rect(80,272,1440,137,CARD)
    s.text(110,294,'两种建模方式',26,TEAL,bold=True)
    s.text(110,346,'因子组合 → 交易信号',29,WHITE,width=600)
    s.text(814,346,'Python 自定义 → 策略逻辑',29,WHITE,width=650)
    s.line(770,310,770,382,LINE,1)
    steps=[('信号与价格','目标权重 / 买卖数量'),('账户模拟','现金 / 持仓 / 交易'),('绩效统计','收益 / 波动 / 回撤'),('账户报告','统计表与图形解释')]
    for i,(t,b) in enumerate(steps):
        x=80+i*369
        s.node(x,493,337,t,b,TEAL if i<2 else BLUE,h=134)
        if i<3:s.line(x+337,560,x+365,560,DIM,2,True)
    s.text(95,708,'交易假设显式输入：成交价格 · 买卖费率 · 交易限制 · 金额限制',28,WHITE,width=1410)
    s.text(95,759,'研究与模拟场景；本页不涉及实盘委托链路。',21,DIM,width=1300)

    s = add('04  /  风险与组合', '把研究观点，转化为受约束的组合决策',
            '风险模型刻画暴露，组合优化在收益、风险与交易成本之间求解。',
            ['Q11','Q12','Q13','Q14','W07','W08','W21'], '核心风险模块包含风险库与多因子风险数据结构，并有 Barra 风格风险模型实现。QSWeb 风险管理更侧重风险库浏览、矩阵展示与风险分解，不将其等同于完整的 Web 风险模型训练平台。组合优化核心和 Web 页面均提供均值方差、风险预算、最大分散化三类目标。核心约束还包括因子暴露、波动与预期收益；Web 服务当前接入预算、权重、换手和非零权重数目约束，具体组合依赖场景及求解器。优化输出不是收益承诺。')
    s.rect(80,274,695,446,CARD)
    s.text(112,302,'风险建模与分析',34,TEAL,width=610,bold=True)
    s.text(112,375,'因子暴露 · 因子协方差 · 特异风险',28,WHITE,width=610)
    s.text(112,457,'Barra 风格风险模型\n风险库读写与风险数据组织\nWeb 矩阵展示与风险分解',27,MUTED,width=610,leading=1.7)
    s.rect(805,274,715,446,CARD)
    s.text(837,302,'组合构建与优化',34,BLUE,width=630,bold=True)
    s.text(837,375,'目标 + 约束 → 求解 → 组合权重',28,WHITE,width=630)
    s.text(837,457,'均值方差 · 风险预算 · 最大分散化\n预算 / 权重 / 换手 / 持仓数等约束\nCVXPY 建模与求解信息',27,MUTED,width=630,leading=1.7)
    s.text(80,758,'核心还支持因子暴露、波动率等约束；Web 接入范围以具体服务实现为准。',22,DIM,width=1440)

    s = add('05  /  研究工程化', 'QSExt 把研究代码组织成可管理的定义与资产',
            '统一定义入口承接研究逻辑，注册中心记录研究对象及其依赖。',
            ['E04','E05','E06','E07','E13'], 'DefModule 是当前统一因子和策略定义框架，支持 defNode、defFactor、defStrategy，运行时上下文管理数据源、日期、截面和模型参数，并可解析依赖。不同项目文档仍可能保留旧 FactorDef/StrategyDef 说明，本介绍采用当前 DefModule。QSRegistry 基于 Neo4j，存储因子、脚本、策略、回测、风险表、优化器及报告等元信息和关系。登记、查询、重建与影响分析提供复用基础；外部代码、数据源和运行环境仍需可用。MCP 工具为助手暴露查询等能力。')
    s.card(80,278,695,353,'DEFMODULE','定义与运行','元信息、模型参数与依赖声明\n统一因子 / 策略 / 混合脚本\n运行配置与数据源管理',TEAL)
    s.card(805,278,715,353,'QSREGISTRY','登记与复用','保存研究对象与依赖关系\n检索、血缘、影响分析与重建\n通过 MCP 提供工具访问',BLUE)
    s.line(775,452,805,452,DIM,2,True)
    s.text(88,700,'研究脚本',28,WHITE,width=240)
    s.line(310,718,390,718,TEAL,2,True)
    s.text(422,700,'统一定义',28,WHITE,width=260)
    s.line(650,718,730,718,TEAL,2,True)
    s.text(762,700,'计算与结果',28,WHITE,width=280)
    s.line(1010,718,1090,718,TEAL,2,True)
    s.text(1122,700,'研究资产登记',28,WHITE,width=380)

    s = add('06  /  统一工作台', 'QSWeb 让研究链路在一张画布中可见',
            '来源、因子加工、策略与回测输出，可以在同一计算图中组织。',
            ['W01','W15','W16','W17'], '本页为真实 QSWeb 前端编辑器的本地截图。通过项目已有浏览器测试挂载入口加载真实组件，接口由确定性演示夹具替代；图结构用于展示连接关系，没有访问真实业务数据库，也没有计算策略绩效。截图中的组件名称为演示定义，图处于待执行状态。页面提供组件库、节点设置、运行及诊断与结果入口，实际功能由 QSExt Graph 与后端服务支持。')
    screenshot=out/'assets/qsweb_graph.png'
    if screenshot.exists():
        s.image(80,266,1180,645,screenshot)
        # 截图裁切在内容区内，保留工作台上方控制区和节点。
        s.items[-1]['crop_bottom'] = 0.17
        s.items[-1]['h'] = 506
    else:
        s.rect(80,266,1180,506,CARD)
        s.text(120,440,'统一计算图工作台',50,WHITE,width=1000)
    for i,(t,b) in enumerate([('组织逻辑','来源与组件连线'),('保存定义','参数与修订留存'),('观察执行','诊断、进度与结果')]):
        s.text(1300,290+i*144,t,29,TEAL,bold=True,width=230)
        s.text(1300,343+i*144,b,23,MUTED,width=220)
    s.text(82,785,'真实前端截图；演示组件与接口夹具，未连接业务数据。',19,DIM,width=1410)

    s = add('06  /  执行链路', '从编辑到运行，关键边界保持清晰',
            '保存图文档、静态校验与实际执行，分别由明确的步骤承接。',
            ['E01','E02','W03','W04','W05','W06'], 'GraphDocument 保存组件定义引用、参数、输入绑定与输出；页面布局属于扩展数据。静态校验不构造业务计算对象，也不访问业务数据库。明确提交任务后，后端读取指定图修订并检查资源，编译构造 Node/Factor，通过 QuantStudio 引擎运行。运行配置区分计算日期与带历史回溯的时点标尺。长任务当前使用内存 asyncio 管理器，WebSocket/轮询提供进度，图运行记录另外持久化；这不是可跨服务重启恢复的分布式任务队列。')
    steps=[('图文档','参数 / 绑定 / 输出'),('静态校验','类型 / 必填 / 依赖'),('目标编译','构造 Node / Factor'),('引擎执行','数据准备与计算'),('结果展示','结果树 / HTML')]
    for i,(t,b) in enumerate(steps):
        x=80+i*294
        s.node(x,329,266,t,b,TEAL if i<2 else BLUE,h=160)
        if i<4:s.line(x+266,410,x+291,410,DIM,2,True)
    s.rect(80,560,706,163,CARD)
    s.text(107,585,'提交时固定执行依据',29,TEAL,bold=True,width=650)
    s.text(107,643,'图修订 · 组件定义 · 授权资源 · 日期与截面',25,WHITE,width=650)
    s.rect(816,560,704,163,CARD)
    s.text(843,585,'运行中提供可见反馈',29,BLUE,bold=True,width=650)
    s.text(843,643,'异步任务 · 节点事件 · 进度推送 · 结果记录',25,WHITE,width=650)

    s = add('07  /  报告与资产', '报告承接计算结果，也保留研究的上下文',
            '场景、布局与渲染分离，让分析输出可以持续复用。',
            ['E07','E08','E09','W09'], 'QSExt ReportGenerator 继承 ReportNode，可参与计算图并被 BTReport 聚合；提供 YAML 配置、布局组件以及 HTML/Markdown 渲染。场景目录包含 single_factor、multi_factor、single_strategy、multi_strategy。图表可使用 Plotly 或 Matplotlib。报告中心提供生成、预览及注册相关流程。图中为能力关系示意，不是实际研究报告截图，也没有虚构业绩图。注册中心可以将报告与因子和回测对象关联，用于检索与追踪。')
    for i,(t,b) in enumerate([('输入结果','因子分析 / 策略回测'),('报告场景','单 / 多因子 · 单 / 多策略'),('呈现方式','图表 / 表格 / 统计摘要')]):
        s.node(80,280+i*145,610,t,b,TEAL if i!=1 else BLUE,h=119)
    s.line(716,488,804,488,DIM,3,True)
    s.rect(840,278,680,435,CARD)
    s.text(876,310,'结构化研究报告',38,WHITE,bold=True,width=600)
    s.line(876,381,1482,381,LINE,1)
    s.text(876,413,'研究对象与参数',29,TEAL,width=600)
    s.text(876,476,'分析结果与图表',29,WHITE,width=600)
    s.text(876,539,'报告文件与关联元信息',29,WHITE,width=600)
    s.text(876,626,'HTML / Markdown',27,MUTED,latin=True,width=600)
    s.text(80,759,'输出既服务于阅读，也能关联回研究对象与计算依据。',29,WHITE,width=1440)

    s = add('07  /  研究扩展', 'AI 与因子挖掘，扩展研究的探索方式',
            '将辅助开发、候选生成与评估，接回已有的研究基础设施。',
            ['E10','E11','E12','E13','W10','W11','W18'], '当前代码具备三个不同层次的能力：AI 工作台使用 Agent SDK 提供多轮辅助开发；GPFactor 通过遗传规划与适应度评价搜索表达式；LLMFactor 包含多轮假设、开发与评价相关流程，QSWeb mining_service 实际包含 llm_factor 提交及执行分支，尽管服务文件顶部仍有仅集成 GP 的旧说明。这里依据实际实现而非旧注释。MCP 可查询注册对象和代码。依赖模型服务、权限、数据和评估配置，不能宣称全自动产出有效 Alpha 或已完成独立有效性验证。')
    s.card(80,284,460,363,'AI WORKBENCH','辅助研究开发','多轮对话与上下文管理\n代码辅助与研究工具调用\n结合 MCP 查询研究资产',TEAL)
    s.card(570,284,460,363,'GPFACTOR','遗传规划搜索','算子与终端因子组合\n候选表达式进化\n使用适应度函数评价',BLUE)
    s.card(1060,284,460,363,'LLMFACTOR','假设驱动探索','生成研究假设\n开发候选因子\n多轮评价与探索记录',GOLD)
    s.text(82,711,'共同基础：数据上下文 · 因子计算 · 回测评估 · 研究结果留存',29,WHITE,width=1430)
    s.text(82,771,'依赖相应模型服务与运行配置；候选有效性仍需独立研究验证。',21,DIM,width=1430)

    s = add('08  /  部署与团队', '面向团队自部署，研究环境与服务目录分离',
            'Web 服务、计算框架、业务数据与用户工作区按职责组织。',
            ['W12','W13','W14','W19','W20'], '系统适合小团队自部署。Web 前端使用 React/TypeScript，后端使用 FastAPI；计算由同环境的 QuantStudio/QSExt 承接。应用数据库保存用户、图修订和运行等信息，Neo4j 负责注册元信息与依赖关系；因子和风险数据存于各自后端，不能混同。配置、共享组件、用户目录和日志从服务目录组织；Notebook/JupyterLab 按部署配置接入。现有任务管理器是内存模型，不承诺高可用、分布式集群或服务重启后的任务恢复。')
    cols=[('访问与交互',['Web 研究工作台','AI / Notebook 入口','JWT 认证与用户管理'],BLUE),('计算与服务',['FastAPI 业务服务','QuantStudio / QSExt','任务进度与结果接口'],TEAL),('数据与工作区',['因子库 / 风险库','注册元信息 / 应用数据','配置 / 组件 / 用户文件'],GOLD)]
    for i,(t,body,c) in enumerate(cols):
        x=80+i*490
        s.rect(x,285,460,350,CARD)
        s.text(x+27,311,t,33,c,bold=True,width=408)
        for j,b in enumerate(body):
            s.dot(x+32,407+j*70,4,c)
            s.text(x+50,389+j*70,b,25,WHITE,width=380)
        if i<2:s.line(x+460,457,x+486,457,DIM,2,True)
    s.rect(80,686,1440,83,CARD2)
    s.text(108,706,'交付形态：本地服务或 Docker Compose；实际能力取决于连接、依赖与组件配置。',26,WHITE,width=1380)

    s = add('09  /  结语', '让研究成果持续积累，让团队能力持续生长',
            'QuantStudio 系统连接数据、计算、验证与知识沉淀。',
            ['Q02','E01','E07','W01'], '让研究成果持续积累，让团队能力持续生长。QuantStudio 系统连接数据、计算、验证与知识沉淀。')

    s = add('附录 A  /  能力边界', '区分核心能力、界面接入与部署前提',
            '能力判断以当前工作区的实际代码为依据。',
            ['Q12','Q14','W01','W06','W07','W08','W10','W14','W21'], '本页供技术评审或问答使用。风险模型实现属于核心，Web 更多是浏览与分析；核心约束比 Web 服务实际接入范围更广，三类优化目标已同时接入前后端；底层有多进程并行，但 Web 任务管理器为内存 asyncio，不能据此宣称分布式任务调度与高可用。注册中心存元信息与关系，数据和文件需独立管理。AI 和外部数据源需依赖服务及授权。历史回测能力不代表已提供实盘交易系统。PPT 未连接业务库或跑策略，因此不作生产可用性或数值结果背书。')
    widths=[260,580,600]; starts=[80,340,920]
    for x,w,t in zip(starts,widths,['主题','代码可确认的能力','需要明确的边界']):
        s.rect(x,268,w,62,CARD2)
        s.text(x+18,283,t,24,TEAL,bold=True,width=w-36)
    table=[('风险与优化','核心风险模型；Web 接入三类优化目标','Web 风险页面侧重浏览与分析'),('图与运行','保存修订、目标编译与运行记录','完整复现需数据、组件及环境版本'),('并行与任务','多进程计算 + 内存异步任务管理','未证明分布式调度、重启恢复与高可用'),('研究资产','Neo4j 元信息、依赖与检索','业务数据和结果文件需独立存储管理'),('AI 与数据源','有助手、挖掘和多种适配器实现','依赖模型服务、数据授权与运行配置'),('策略与交易','历史账户模拟、绩效统计与报告','未在本次考察中验证实盘委托链路')]
    for i,row in enumerate(table):
        y=330+i*71
        for x,w,t in zip(starts,widths,row):
            s.rect(x,y,w,71,CARD if i%2==0 else BG)
            s.text(x+18,y+20,t,22,WHITE if x==80 else MUTED,width=w-36)

    s = add('附录 B  /  资料依据', '以源码和正式文档支撑系统介绍',
            '逐页讲稿与完整来源路径已写入 PowerPoint 备注。',
            ['Q01','Q02','E01','E04','E07','W01','W03'], '资料方法：先阅读项目规范、QuantStudio 通则 Notebook，再检查核心接口与最新实现。QSExt 和 QSWeb 部分旧说明与代码不同，以 Graph、DefModule 和统一计算图工作台的现行实现为主。引用编号在每页页脚展示，完整来源在各页备注与 sources.json 中。sources.json 记录文件 SHA-256、当前 HEAD 以及引用页码；Git HEAD 不能代表未提交工作区状态，文件哈希用于辅助识别本次依据。真实 UI 截图由本地演示夹具生成，未运行真实策略或访问业务数据库。能力介绍不包含生产验收、性能基准与投资业绩结论。')
    data=[('QuantStudio','计算图 · 因子 · 回测 · 风险 · 组合','docs/Core/计算图框架.ipynb\nQuantStudio/Core/Node.py\nQuantStudio/BackTest/SectionFactor/api.py'),
          ('QSExt','图协议 · 定义框架 · 注册 · 报告 · 挖掘','docs/Graph/README.md\ndocs/DefModule/DefModule.md\nQSExt/QSRegistry/QSGraphDB.py'),
          ('QSWeb','统一工作台 · 执行服务 · 用户与任务','docs/13-统一计算图工作台.md\nbackend/app/services/graph_service.py\nfrontend/src/App.tsx')]
    for i,(t,b,paths) in enumerate(data):
        y=275+i*148
        s.rect(80,y,1440,128,CARD)
        s.text(106,y+22,t,32,[GOLD,TEAL,BLUE][i],latin=True,bold=True,width=300)
        s.text(415,y+20,b,24,WHITE,width=1060)
        s.text(415,y+65,paths,17,MUTED,width=1060,leading=1.3)
    s.text(82,745,'考察日期：2026.10.20  /  本地工作区快照，包含未提交改动。',23,WHITE,width=1420)
    s.text(82,788,'Barra 示例取自 QSResearch 的两份定义脚本，完整路径与处理顺序见第 8 页备注。',19,DIM,width=1420)

    for i,page in enumerate(pages,1):page.footer(i,len(pages))
    return pages


def render_pptx(pages, dest):
    """将场景写入可编辑的 PowerPoint。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = 'QuantStudio 系统介绍'
    prs.core_properties.subject = 'QuantStudio 系统介绍'
    prs.core_properties.author = 'QuantStudio'
    prs.core_properties.keywords = '量化研究,计算图,因子,回测,风险,组合,研究资产'
    scale = 914400 / 120
    def emu(v):return int(v * scale)
    for p in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*rgb(BG))
        for a in p.items:
            k=a['kind']
            if k=='text':
                shp=slide.shapes.add_textbox(emu(a['x']),emu(a['y']),emu(a['w']),emu(a['h']))
                tf=shp.text_frame
                tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                para=tf.paragraphs[0]
                para.space_before=para.space_after=Pt(0)
                para.line_spacing=1.0
                r=para.add_run();r.text=a['text']
                r.font.name='Segoe UI' if a['latin'] else 'Microsoft YaHei'
                r.font.size=Pt(a['size']*0.6);r.font.bold=a['bold'];r.font.color.rgb=RGBColor(*rgb(a['color']))
                rpr=r._r.get_or_add_rPr()
                for tag in ['a:ea','a:cs']:
                    el=OxmlElement(tag);el.set('typeface','Microsoft YaHei');rpr.append(el)
                para._p.get_or_add_pPr().set('fontAlgn','base')
            elif k in ('rect','circle'):
                shape=MSO_SHAPE.OVAL if k=='circle' else MSO_SHAPE.ROUNDED_RECTANGLE if a.get('radius') else MSO_SHAPE.RECTANGLE
                shp=slide.shapes.add_shape(shape,emu(a['x']),emu(a['y']),emu(a['w']),emu(a['h']))
                if a.get('radius'):
                    shp.adjustments[0]=0.06
                shp.fill.solid();shp.fill.fore_color.rgb=RGBColor(*rgb(a['fill']))
                if a.get('stroke'):
                    shp.line.color.rgb=RGBColor(*rgb(a['stroke']));shp.line.width=Pt(0.6)
                else:shp.line.fill.background()
            elif k=='line':
                shp=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,emu(a['x']),emu(a['y']),emu(a['x2']),emu(a['y2']))
                shp.line.color.rgb=RGBColor(*rgb(a['color']));shp.line.width=Pt(a['width']*0.6)
                if a.get('arrow'):
                    end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');shp.line._get_or_add_ln().append(end)
            elif k=='image':
                shp=slide.shapes.add_picture(a['path'],emu(a['x']),emu(a['y']),emu(a['w']),emu(a['h']))
                shp.crop_bottom=a.get('crop_bottom',0)
        refs='\n'.join(f'[{key}] {SOURCES[key][0]}/{SOURCES[key][1]}' for key in p.sources)
        slide.notes_slide.notes_text_frame.text = f'【讲稿】\n{p.notes}\n\n【来源】\n{refs}\n\n【范围】依据 2026-10-20 本地工作区；涉及推导、场景示意与采用建议时已说明。'
    prs.save(dest)


def render_pdf_and_images(pages, out):
    """按相同场景生成矢量 PDF 与逐页图像。"""
    pdf=fitz.open()
    preview=out/'preview';preview.mkdir(exist_ok=True)
    # 仅清理本生成器遗留的多余页图，避免删页后仍显示旧预览。
    for old in preview.glob('[0-9][0-9].png'):
        if int(old.stem)>len(pages):
            old.unlink()
    fonts={}
    for idx,p in enumerate(pages,1):
        page=pdf.new_page(width=960,height=540)
        for name,file in [('cn',FONT_FILE),('cnb',BOLD_FILE),('en',LATIN_FILE),('enb',LATIN_BOLD)]:
            page.insert_font(fontname=name,fontfile=file)
        for a in p.items:
            k=a['kind'];sc=0.6
            if k=='text':
                fn=('en' if a['latin'] else 'cn')+('b' if a['bold'] else '')
                if fn not in fonts:
                    fonts[fn]=fitz.Font(fontfile={'cn':FONT_FILE,'cnb':BOLD_FILE,'en':LATIN_FILE,'enb':LATIN_BOLD}[fn])
                # 字体顶端对齐，与 PPT 文本框使用统一的留白。
                baseline=(a['y']+a['size']*1.05)*sc
                page.insert_text((a['x']*sc,baseline),a['text'],fontname=fn,fontsize=a['size']*sc,color=tuple(v/255 for v in rgb(a['color'])))
            elif k in ('rect','circle'):
                rect=fitz.Rect(a['x']*sc,a['y']*sc,(a['x']+a['w'])*sc,(a['y']+a['h'])*sc)
                fill=tuple(v/255 for v in rgb(a['fill']))
                color=tuple(v/255 for v in rgb(a['stroke'])) if a.get('stroke') else None
                if k=='circle':page.draw_oval(rect,color=color,fill=fill,width=0.6)
                else:page.draw_rect(rect,color=color,fill=fill,width=0.6)
            elif k=='line':
                x,y,x2,y2=[a[t]*sc for t in ['x','y','x2','y2']]
                color=tuple(v/255 for v in rgb(a['color']))
                page.draw_line((x,y),(x2,y2),color=color,width=a['width']*sc)
                if a.get('arrow'):
                    angle=math.atan2(y2-y,x2-x)
                    points=[(x2,y2)]+[(x2-7*math.cos(angle+d),y2-7*math.sin(angle+d)) for d in [-0.45,0.45]]
                    shape=page.new_shape();shape.draw_polyline(points+[points[0]]);shape.finish(color=color,fill=color);shape.commit()
            elif k=='image':
                rect=fitz.Rect(a['x']*sc,a['y']*sc,(a['x']+a['w'])*sc,(a['y']+a['h'])*sc)
                img=Image.open(a['path']).convert('RGB')
                if a.get('crop_bottom'):img=img.crop((0,0,img.width,round(img.height*(1-a['crop_bottom']))))
                import io
                b=io.BytesIO();img.save(b,format='PNG')
                page.insert_image(rect,stream=b.getvalue(),keep_proportion=False)
        pix=page.get_pixmap(matrix=fitz.Matrix(1600/960,1600/960),alpha=False)
        pix.save(preview/f'{idx:02d}.png')
    pdf.subset_fonts()
    pdf.save(out/'QuantStudio_系统介绍.pdf',garbage=4,deflate=True)
    pdf.close()


def validate(pages, out):
    """检查源文件、文本边界、页数和可读取性。"""
    issues=[]
    source_rows={}
    for key,(repo,rel) in SOURCES.items():
        path=ROOT.parent/repo/rel
        if not path.is_file():
            issues.append(f'缺失来源：{key} {repo}/{rel}')
        source_rows[key]={'project':repo,'path':rel,'sha256':hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None,
                          'slides':[i for i,p in enumerate(pages,1) if key in p.sources]}
    for idx,p in enumerate(pages,1):
        for a in p.items:
            if a['kind']!='text':continue
            if a['x']<0 or a['y']<0 or a['x']+a['w']>W+4 or a['y']+a['h']>H+4:
                issues.append(f'第 {idx} 页文本越界：{a["text"]}')
            if 810<a['y']<827:
                issues.append(f'第 {idx} 页正文进入页脚：{a["text"]}')
    prs=Presentation(out/'QuantStudio_系统介绍.pptx')
    if len(prs.slides)!=len(pages):issues.append('PPT 页数不一致')
    doc=fitz.open(out/'QuantStudio_系统介绍.pdf')
    for i,p in enumerate(doc):
        if len(p.get_text().strip())<50:issues.append(f'第 {i+1} 页 PDF 文本异常')
    doc.close()
    heads={repo:subprocess.run(['git','-c',f'safe.directory={(ROOT.parent/repo).as_posix()}','rev-parse','HEAD'],cwd=ROOT.parent/repo,capture_output=True,text=True,check=True).stdout.strip() for repo in sorted({repo for repo,_ in SOURCES.values()})}
    (out/'sources.json').write_text(json.dumps({'date':'2026-10-20','scope':'本地工作区，包含未提交改动；文件哈希标识本次引用内容。','git_heads':heads, 'sources':source_rows},ensure_ascii=False,indent=2),encoding='utf-8')
    result={'slides':len(pages),'issues':issues,'checks':['来源文件存在','文本画布边界','PPTX 可重开','PDF 文本可提取','讲稿和来源备注']}
    (out/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def make_contact_sheet(pages, out):
    """生成带页码的全稿总览。"""
    cols=4;cw=400;ch=250
    im=Image.new('RGB',(cols*cw,math.ceil(len(pages)/cols)*ch),rgb(BG));draw=ImageDraw.Draw(im)
    for i,p in enumerate(pages):
        thumb=Image.open(out/'preview'/f'{i+1:02d}.png').convert('RGB').resize((384,216),Image.Resampling.LANCZOS)
        x=(i%cols)*cw+8;y=(i//cols)*ch+6
        im.paste(thumb,(x,y));draw.text((x+5,y+220),f'{i+1:02d}  {p.title or "QuantStudio 系统介绍"}',font=font(13),fill=rgb(MUTED))
    im.save(out/'总览.jpg',quality=94)


def main():
    """创建演示文件并校验。"""
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'docs/presentation')
    args=parser.parse_args();out=args.output.resolve();out.mkdir(parents=True,exist_ok=True)
    pages=deck(out)
    render_pptx(pages,out/'QuantStudio_系统介绍.pptx')
    render_pdf_and_images(pages,out)
    make_contact_sheet(pages,out)
    result=validate(pages,out)
    (out/'slides.json').write_text(json.dumps([{'page':i,'title':p.title or 'QuantStudio 系统介绍','chapter':p.chapter,'notes':p.notes,'sources':p.sources} for i,p in enumerate(pages,1)],ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if result['issues']:raise SystemExit(1)


if __name__=='__main__':
    main()
